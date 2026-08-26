#!/usr/bin/env python3
"""生成知华科技 ZhuaTech VMS 面向最终用户的产品使用手册。

Copyright 2026 上海如静知华信息科技有限公司 · https://www.zhuatech.cn/
"""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "知华科技-ZhuaTech-VMS-产品使用手册-V1.2.docx"
IMAGES = ROOT / "docs" / "images"

GREEN = "176B57"
DARK = "173C33"
MINT = "E7F3EF"
PALE = "F3F7F5"
GOLD = "D5A044"
RED = "A64235"
INK = "18221F"
MUTED = "65736E"
WHITE = "FFFFFF"
LINE = "DCE4E0"
FONT_CN = "PingFang SC"


def font(run, size=11, bold=False, color=INK, name=FONT_CN):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url, color=GREEN):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    font(run, 8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    font(paragraph.add_run(" 页"), 8.5, color=MUTED)


def set_alt(shape, description):
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", description)


def p(doc, text="", size=11, color=INK, bold=False, align=None, before=0, after=6, line=1.25, keep=False):
    para = doc.add_paragraph()
    if text:
        font(para.add_run(text), size, bold, color)
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = line
    para.paragraph_format.keep_with_next = keep
    if align is not None:
        para.alignment = align
    return para


def heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    font(para.add_run(text), 16 if level == 1 else 13, True, GREEN if level == 1 else DARK)
    para.paragraph_format.space_before = Pt(18 if level == 1 else 14)
    para.paragraph_format.space_after = Pt(10 if level == 1 else 7)
    para.paragraph_format.keep_with_next = True
    return para


def kicker(doc, text):
    return p(doc, text.upper(), 8.5, GREEN, True, before=0, after=4, keep=True)


def callout(doc, title, body, fill=MINT, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(3)
    font(cp.add_run(title), 10.5, True, accent)
    bp = cell.add_paragraph()
    bp.paragraph_format.space_after = Pt(0)
    bp.paragraph_format.line_spacing = 1.2
    font(bp.add_run(body), 9.5, False, INK)
    p(doc, "", after=4)


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    font(para.add_run(text), 10.5, False, INK)
    para.paragraph_format.left_indent = Inches(.375)
    para.paragraph_format.first_line_indent = Inches(-.188)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.25
    return para


def step(doc, number, title, body):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [720, 8640])
    left, right = table.rows[0].cells
    shade(left, GREEN)
    shade(right, PALE)
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(lp.add_run(str(number)), 13, True, WHITE)
    rp = right.paragraphs[0]
    font(rp.add_run(title), 10.5, True, DARK)
    rp2 = right.add_paragraph()
    rp2.paragraph_format.space_after = Pt(0)
    font(rp2.add_run(body), 9.5, False, MUTED)
    p(doc, "", after=2)


def image(doc, filename, caption, width=6.2):
    path = IMAGES / filename
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    shape = para.add_run().add_picture(str(path), width=Inches(width))
    set_alt(shape, caption)
    cap = p(doc, caption, 8.5, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=8, keep=True)
    cap.style = doc.styles["Caption"]


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_table_geometry(t, widths)
    set_repeat_header(t.rows[0])
    for i, text in enumerate(headers):
        shade(t.rows[0].cells[i], DARK)
        para = t.rows[0].cells[i].paragraphs[0]
        font(para.add_run(text), 9, True, WHITE)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            if len(t.rows) % 2 == 1:
                shade(cells[i], PALE)
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.15
            font(para.add_run(str(text)), 9, False, INK)
    return t


def new_page(doc):
    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.42)
    section.footer_distance = Inches(.42)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name in ("Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    doc.core_properties.title = "知华科技 ZhuaTech VMS 产品使用手册 V1.2"
    doc.core_properties.subject = "访客预约、审批、签到、通行、离场、风险预警与资源配置"
    doc.core_properties.author = "上海如静知华信息科技有限公司"
    doc.core_properties.keywords = "知华科技,VMS,访客管理系统,访客预约,访客通行码,园区门禁"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font(header.add_run("知华科技 · ZhuaTech VMS  产品使用手册"), 8.5, True, MUTED)
    add_page_number(section.footer.paragraphs[0])

    # Cover - editorial_cover pattern with branded restraint.
    p(doc, "ZHUATECH  ·  PRODUCT GUIDE", 9, GREEN, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=72, after=22)
    p(doc, "知华科技", 15, DARK, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p(doc, "ZhuaTech VMS", 31, DARK, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=7)
    p(doc, "访客预约与通行管理平台", 18, GREEN, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    p(doc, "产品使用手册", 13, MUTED, False, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    callout(doc, "从预约到离场的可追踪闭环", "适用于园区前台、行政接待、内部接待人、安保人员和系统管理员。本文档以社区源码版 V1.4 的实际页面与已实现功能为准。")
    p(doc, "手册版本  V1.2  ·  产品版本  V1.4", 10, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=42, after=4)
    p(doc, "发布日期  2026 年 8 月", 10, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    p(doc, "上海如静知华信息科技有限公司", 10.5, DARK, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=16, after=5)
    link_p = p(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_hyperlink(link_p, "https://www.zhuatech.cn/", "https://www.zhuatech.cn/")

    new_page(doc)
    kicker(doc, "READ THIS FIRST")
    heading(doc, "开始使用前", 1)
    p(doc, "本手册面向日常使用人员，重点说明每个模块解决什么问题、由谁使用、如何操作，以及不同状态代表什么。技术接口和二次开发细节请参见项目 docs/API.md 与 docs/ARCHITECTURE.md。")
    callout(doc, "开源版本使用边界", "本工程仅能用于个人非商业学习、研究和技术交流。未经上海如静知华信息科技有限公司书面授权，不得商用、对外收费部署、SaaS 运营、外包交付或投标。", "FFF4E1", GOLD)
    heading(doc, "角色与工作入口", 2)
    table(doc, ["角色", "主要入口", "典型工作"], [
        ["系统管理员", "管理端", "资源维护、基础设置、运营报表与审计查看"],
        ["前台 / 行政", "管理端", "新建预约、资料核对、签到与离场确认"],
        ["内部接待人", "移动端", "预约审批、接待提醒、访客状态跟进"],
        ["安保人员", "管理端 / 移动端", "身份复核、通行异常、黑名单和应急清点"],
        ["访客", "移动端", "提交信息、展示通行码、配合签到与签退"],
    ], [1500, 1700, 6160])
    heading(doc, "章节导航", 2)
    table(doc, ["模块", "本手册内容"], [
        ["运营总览", "指标、流程进度、现场接待状态与最近预约"],
        ["业务协同", "预约新增、审批、驳回、签到、通行与离场"],
        ["资源中心", "访客档案、黑名单、接待人、区域与门禁点"],
        ["风险预警", "准入判断、异常上报、处置留痕与应急清点"],
        ["审计报表", "预约结构、档案风险、资源可用率和操作日志"],
        ["基础设置", "园区、审批、通行、通知、数据留存与账号权限"],
        ["移动工作台", "预约、签到、通行码、待办、消息与异常上报"],
    ], [1900, 7460])

    new_page(doc)
    kicker(doc, "END-TO-END FLOW")
    heading(doc, "1. 来访业务全流程", 1)
    p(doc, "ZhuaTech VMS 用一条状态链连接访客、接待人、前台和安保。每次关键操作都会改变预约状态，操作者可据此判断下一步。")
    for idx, item in enumerate([
        ("提交预约", "填写访客、单位、手机、接待人、到访时间、访问区域和事由。"),
        ("接待审批", "接待人确认来访目的；受限区域或风险关注记录由安保复核。"),
        ("身份核验", "到访时核对预约编号、证件信息和人员名单。"),
        ("签发通行码", "审批通过后生成一次性数字通行码，并绑定区域与有效时段。"),
        ("访客签到", "核验通过后确认入园，状态变为“已到访”。"),
        ("离场确认", "访客离场后签退，系统回收通行权限并保留记录。"),
    ], 1):
        step(doc, idx, item[0], item[1])
    heading(doc, "状态快速判断", 2)
    table(doc, ["当前状态", "说明", "下一步"], [
        ["待审批", "预约已提交，尚未同意", "通过并发码 / 驳回 / 取消"],
        ["已审批", "接待已确认，通行码可用", "身份核验后签到"],
        ["已到访", "访客当前在园", "关注超时与异常，离场时签退"],
        ["已离场", "访问结束，权限已回收", "归档或查询历史记录"],
        ["已驳回 / 已取消", "预约不再生效", "补充资料后重新发起"],
    ], [1500, 3850, 4010])

    new_page(doc)
    kicker(doc, "MODULE 01")
    heading(doc, "2. 运营总览", 1)
    p(doc, "运营总览是管理员、前台与安保人员进入系统后的第一工作面。它用于快速回答：今天有多少预约、还有多少待审批、当前多少人在园、有哪些预警需要处理。")
    image(doc, "vms-reception-dashboard.png", "图 2-1  管理端运营总览：指标、流程进度、健康度与最近预约", 6.35)
    heading(doc, "主要信息", 2)
    bullet(doc, "今日预约：当前预约总量，用于判断前台接待压力。")
    bullet(doc, "待审批：尚未完成接待确认的预约，应优先处理。")
    bullet(doc, "当前在园：已经签到但尚未离场的访客人数。")
    bullet(doc, "待处理预警：超时、资料缺失、通行异常等未闭环事项。")
    callout(doc, "推荐工作习惯", "每天上班后先查看待审批和预警，下班前核对“当前在园”是否仍有访客未签退。")

    new_page(doc)
    kicker(doc, "MODULE 02")
    heading(doc, "3. 业务协同：预约管理", 1)
    p(doc, "业务协同集中展示预约编号、访客信息、来访时间、接待人、访问区域、风险等级和当前状态。支持关键词搜索和状态筛选。")
    image(doc, "vms-appointment-workflow.png", "图 3-1  预约与接待协同列表：风险和状态一屏可见", 4.1)
    heading(doc, "新建预约", 2)
    for idx, item in enumerate([
        ("点击“新建预约”", "管理端右上角和移动端快捷服务均可进入。"),
        ("填写访客资料", "至少填写姓名、手机、单位、接待人和来访事由。"),
        ("安排访问", "选择日期、时段、访问区域和来访人数。"),
        ("提交", "系统生成预约编号，初始状态为“待审批”。"),
    ], 1):
        step(doc, idx, item[0], item[1])
    callout(doc, "风险提醒", "受限区域、多人、夜间或身份待核验会触发复核；黑名单命中应停止预约。", "FCEDE9", RED)

    kicker(doc, "WORKFLOW ACTIONS")
    heading(doc, "4. 审批、签到、通行与离场", 1)
    heading(doc, "4.1 审批或驳回", 2)
    p(doc, "在预约列表点击“处理”进入详情。确认访客、接待人、时间、区域和来访事由后，选择“通过并发码”或“驳回”。审批通过后系统签发六位数字通行码。")
    table(doc, ["操作", "允许状态", "结果"], [
        ["通过并发码", "待审批", "状态变为已审批，并生成通行码"],
        ["驳回", "待审批", "状态变为已驳回，不允许签到"],
        ["取消预约", "未离场且未取消", "状态变为已取消，通行码失效"],
        ["确认签到", "已审批", "状态变为已到访，计入在园人数"],
        ["确认离场", "已到访", "状态变为已离场，回收通行权限"],
    ], [1800, 2200, 5360])
    heading(doc, "4.2 到访核验", 2)
    bullet(doc, "输入预约编号或六位通行码，系统返回凭证是否有效及对应访客信息。")
    bullet(doc, "核对访问区域和有效时段；受限区域由安保人员再次确认。")
    bullet(doc, "核验完成后点击“确认签到”，不得在未审批状态下强制签到。")
    heading(doc, "4.3 离场签退", 2)
    p(doc, "访客离开园区时由前台或接待人确认离场。签退后系统不再将其计入在园人数；如超过预约时段仍未签退，应在风险预警中跟进。")
    callout(doc, "系统校验", "后端会校验状态顺序。例如“待审批”不能直接签到，“已审批”不能直接离场，错误操作返回冲突提示。")

    new_page(doc)
    kicker(doc, "MODULE 03")
    heading(doc, "5. 资源中心", 1)
    p(doc, "资源中心分为访客档案和接待人与区域两类信息。日常人员通常只查询；黑名单维护和基础资源变更应由授权管理员执行。")
    image(doc, "vms-resource-center.png", "图 5-1  访客档案：身份核验、来访次数与黑名单状态", 6.0)
    heading(doc, "5.1 访客档案与黑名单", 2)
    bullet(doc, "档案显示访客单位、脱敏手机、身份核验状态、累计来访次数和最近到访日期。")
    bullet(doc, "前台核验证件原件后可标记“身份核验通过”；撤销核验同样会进入操作日志。")
    bullet(doc, "加入黑名单后，后续预约应触发拦截或人工复核。解除前需要记录复核原因。")
    heading(doc, "5.2 接待人、区域与门禁点", 2)
    table(doc, ["资源类型", "用途", "示例状态"], [
        ["接待人", "绑定内部责任人和所属部门", "启用"],
        ["访问区域", "限制访客可进入的空间范围", "开放 / 审批开放"],
        ["门禁点", "标记闸机或前台核验点", "在线 / 离线"],
    ], [1700, 5400, 2260])
    p(doc, "管理员可以新增、编辑和删除资源。资源编码必须唯一；停用接待人、关闭区域或将门禁点设为离线后，应同步检查未结束预约。")

    new_page(doc)
    kicker(doc, "MODULE 04")
    heading(doc, "6. 风险预警与异常处置", 1)
    p(doc, "风险预警用于集中处理超时未离场、资料待补充、通行异常、证件异常和现场异常。每条记录包含关联预约、等级、责任人、状态和处置结果。")
    image(doc, "vms-risk-alerts.png", "图 6-1  风险预警：待处理、高风险、闭环数和异常明细", 4.0)
    heading(doc, "处置步骤", 2)
    for idx, item in enumerate([
        ("判断等级", "高风险优先处理；中风险应在当班内闭环。"),
        ("核对关联预约", "确认访客、接待人、最后状态和访问区域。"),
        ("线下处置", "联系接待人、前台或安保人员核实现场情况。"),
        ("完成处置", "填写处置结果并点击完成，系统保留操作留痕。"),
    ], 1):
        step(doc, idx, item[0], item[1])
    callout(doc, "应急疏散访客清点", "紧急时按入场、离场、集合点签到和接待确认计算未清点访客，并结合门禁位置现场查找。", "FCEDE9", RED)

    kicker(doc, "MODULE 05")
    heading(doc, "7. 运营报表与操作审计", 1)
    p(doc, "审计报表面向系统管理员，用于查看预约状态结构、访客档案风险、园区资源可用情况及最近 100 条关键操作。")
    image(doc, "vms-audit-report.png", "图 7-1  运营报表与审计日志：指标、状态分布和操作责任人", 3.8)
    heading(doc, "7.1 运营报表", 2)
    bullet(doc, "预约状态分布用于识别待审批积压、当前在园和未完成离场记录。")
    bullet(doc, "档案风险显示黑名单和身份待核验人数，资源指标显示接待人、区域和门禁可用情况。")
    heading(doc, "7.2 操作审计", 2)
    table(doc, ["字段", "用户如何理解"], [
        ["模块 / 操作", "预约、核验、资源、预警或设置变更"],
        ["业务编号", "关联预约、资源或预警编号"],
        ["操作人", "执行操作的登录账号"],
        ["说明", "操作结果或填写的业务备注"],
    ], [2300, 7060])
    callout(doc, "审计使用建议", "定期查看黑名单、身份核验、资源删除、预约状态变更和风险处置记录。审计日志用于追溯，不应由普通运营账号删除。")

    kicker(doc, "MODULE 06")
    heading(doc, "8. 基础设置与权限", 1)
    p(doc, "基础设置决定预约审批方式、通行码有效范围、消息渠道和数据留存期限，仅系统管理员可修改；保存后写入数据库，服务重启后仍然生效。")
    image(doc, "vms-settings.png", "图 8-1  基础设置：园区审批、数据通知和演示账号", 6.35)
    table(doc, ["设置项", "作用", "社区源码版默认值"], [
        ["园区名称", "显示当前管理范围", "上海创新园区"],
        ["审批模式", "决定预约由谁批准", "接待人审批 + 安保复核"],
        ["通行码有效范围", "限制可通行时段", "预约时段前后 30 分钟"],
        ["访客记录保留天数", "控制隐私数据保存周期", "180 天"],
        ["通知渠道", "发送待办与提醒", "站内消息"],
    ], [2000, 3800, 3560])
    callout(doc, "账号说明", "演示账号 admin 拥有全部模块权限；operator 可处理预约、访客与预警，但不能访问管理员设置。上线前必须更换默认密码。")

    new_page(doc)
    kicker(doc, "MOBILE WORKSPACE")
    heading(doc, "9. 移动工作台", 1)
    p(doc, "移动端面向接待人、前台和访客。首页聚合今日接待任务、待审批、提醒和风险关注，并提供四个高频入口。")
    t = doc.add_table(rows=1, cols=2)
    set_table_geometry(t, [4680, 4680])
    for cell, filename, caption in zip(t.rows[0].cells,
        ["vms-mobile-workbench.png", "vms-mobile-pass.png"],
        ["移动工作台首页", "访客通行码"]):
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = para.add_run().add_picture(str(IMAGES / filename), width=Inches(2.15))
        set_alt(shape, caption)
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(cp.add_run(caption), 8.5, True, MUTED)
    heading(doc, "快捷服务", 2)
    table(doc, ["入口", "用户可以完成"], [
        ["发起预约", "填写访客、接待、时间、区域和事由并提交"],
        ["访客签到", "输入预约编号或六位通行码，校验通行凭证"],
        ["通行码", "查看六位通行码、有效时间、访问区域和接待人"],
        ["异常上报", "记录异常类型、等级、关联预约和责任人"],
        ["事项", "查看待审批、已审批、已到访等接待任务"],
        ["消息", "查看接待提醒和风险通知"],
    ], [1800, 7560])

    new_page(doc)
    kicker(doc, "MOBILE PROCEDURES")
    heading(doc, "10. 移动端操作说明", 1)
    heading(doc, "10.1 接待人审批", 2)
    for idx, item in enumerate([
        ("打开待办", "首页点击待办事项或底部“事项”。"),
        ("查看详情", "核对访客、时间、区域、人数和风险提示。"),
        ("作出决定", "选择通过审批、驳回或返回补充资料。"),
    ], 1):
        step(doc, idx, item[0], item[1])
    heading(doc, "10.2 展示通行码", 2)
    p(doc, "审批通过后进入“通行码”。页面显示数字通行码、有效日期、有效时段、访问区域和接待人。通行码仅限本人使用，不应截屏转发。")
    heading(doc, "10.3 异常上报", 2)
    p(doc, "选择异常类型和风险等级，说明时间、地点和具体情况，尽量关联预约编号。提交后由安保中心或指定责任人处理。")
    heading(doc, "10.4 到访与离场", 2)
    bullet(doc, "到访：前台核验身份和预约后确认签到。")
    bullet(doc, "在园：访客按授权区域和时段通行，接待人持续负责。")
    bullet(doc, "离场：由前台或接待人确认签退，通行权限随即回收。")

    new_page(doc)
    kicker(doc, "RISK RULES")
    heading(doc, "11. 准入风险判断", 1)
    p(doc, "社区源码版内置可解释的规则示例，用于帮助学习如何把业务条件转换为准入决定。它不是专业安防产品，不替代人工核验。")
    table(doc, ["风险条件", "系统处理建议"], [
        ["黑名单命中", "直接 REJECT，停止预约并交由安保复核"],
        ["夜间访问", "提高风险分；配置夜间通行时段并通知值班安保"],
        ["受限区域", "提高风险分；限定门禁权限并要求陪同"],
        ["接待人未确认", "进入人工复核；联系内部接待人"],
        ["身份未核验", "进入人工复核；到访前完成实名与证件核验"],
        ["普通预约且资料完整", "APPROVE；按预约时段签发一次性通行码"],
    ], [2500, 6860])
    heading(doc, "输出结果", 2)
    table(doc, ["决定", "含义"], [
        ["APPROVE", "风险较低，可以按预约规则继续"],
        ["MANUAL_REVIEW", "需要接待人或安保人工复核"],
        ["REJECT", "高风险或黑名单命中，不允许继续"],
    ], [2200, 7160])
    callout(doc, "外部能力预留", "身份证件核验、闸机门禁、访客机、企业微信和短信均为预留对接能力，社区源码版未集成真实第三方设备或商业服务。")

    new_page(doc)
    kicker(doc, "DATA & PRIVACY")
    heading(doc, "12. 数据、隐私与安全", 1)
    heading(doc, "12.1 建议收集的信息", 2)
    p(doc, "仅收集完成访问管理所必需的信息。手机号在列表中默认脱敏展示；证件照片、生物特征等敏感数据不应在未具备合规条件时采集。")
    table(doc, ["数据类型", "用途", "建议保护措施"], [
        ["姓名、手机、单位", "预约联系与身份核对", "脱敏展示、权限控制、到期删除"],
        ["访问时间与区域", "通行授权和事件追溯", "按岗位授权、保留审计记录"],
        ["状态与异常记录", "接待闭环和风险处置", "限定管理员和安保人员访问"],
        ["证件 / 照片（预留）", "实名或人证核验", "加密存储、单独授权、短期保存"],
    ], [2200, 3000, 4160])
    heading(doc, "12.2 上线检查", 2)
    bullet(doc, "更换 admin、operator 的默认密码，不对公网暴露演示账号。")
    bullet(doc, "启用 HTTPS，限制数据库和后端端口的网络访问范围。")
    bullet(doc, "根据本组织隐私政策配置保存期限和删除机制。")
    bullet(doc, "接入真实门禁、短信、企业微信前完成供应商安全评估。")
    bullet(doc, "定期检查超时未离场、黑名单变更和风险处置记录。")

    new_page(doc)
    kicker(doc, "FAQ")
    heading(doc, "13. 常见问题", 1)
    faqs = [
        ("为什么预约不能直接签到？", "预约必须先处于“已审批”状态。请由接待人确认来访目的；受限区域还需安保复核。"),
        ("审批通过后在哪里查看通行码？", "移动端点击“通行码”。管理端预约详情也会显示已签发的六位数字码。"),
        ("访客已经离开但仍显示在园怎么办？", "在预约详情中执行“确认离场”。若已超时，可在风险预警中先完成现场核对再签退。"),
        ("如何处理黑名单访客？", "在资源中心确认黑名单状态，不应继续审批；由安保人员核实原因后决定是否解除。"),
        ("企业微信、短信和闸机可以直接用吗？", "社区源码版只预留对接方式，没有真实对接。需要自行开发配置或联系知华科技授权定制。"),
        ("为什么 operator 看不到基础设置？", "基础设置属于管理员权限。operator 仅用于日常预约协同和风险处置。"),
        ("在哪里查看谁修改了资源或预约？", "管理员进入“审计报表”，刷新后可按时间查看模块、操作、业务编号、操作人和说明。"),
        ("可以用于企业正式运营吗？", "未经上海如静知华信息科技有限公司书面授权不得商用或生产部署。")
    ]
    for question, answer in faqs:
        p(doc, question, 11, DARK, True, before=7, after=3, keep=True)
        p(doc, answer, 10.5, MUTED, after=7)

    new_page(doc)
    kicker(doc, "INSTALLATION")
    heading(doc, "14. 安装与体验", 1)
    p(doc, "最简单的体验方式是使用 Docker Compose 启动 MySQL、Java 后端和前端。")
    heading(doc, "环境要求", 2)
    table(doc, ["组件", "建议版本"], [
        ["Docker / Compose", "支持 compose.yaml 的当前稳定版本"],
        ["Java", "21（本地开发）"],
        ["Node.js", "22（本地开发）"],
        ["MySQL", "8.4"],
    ], [2500, 6860])
    heading(doc, "启动步骤", 2)
    for idx, item in enumerate([
        ("准备配置", "复制 .env.example 为 .env，并修改数据库和演示账号密码。"),
        ("启动服务", "在项目根目录执行 docker compose up --build。"),
        ("打开系统", "浏览器访问 http://localhost:8090。"),
        ("登录体验", "演示管理员 admin / admin123；运营人员 operator / operator123。"),
    ], 1):
        step(doc, idx, item[0], item[1])
    callout(doc, "重要", "默认账号和密码仅供本地学习演示，严禁直接用于互联网或生产环境。", "FCEDE9", RED)

    new_page(doc)
    kicker(doc, "SUPPORT")
    heading(doc, "15. 许可边界与技术支持", 1)
    heading(doc, "15.1 非商业使用说明", 2)
    p(doc, "本工程仅能用于个人非商业学习、研究和技术交流。未经上海如静知华信息科技有限公司书面授权，不得用于企业内部生产、商业部署、SaaS、收费下载、售卖、外包交付、投标、品牌替换或任何直接、间接商业用途。具体以项目 LICENSE 为准。")
    heading(doc, "15.2 深度开发与商业授权", 2)
    p(doc, "如需接入身份证件核验、闸机门禁、访客机、企业微信、短信、多园区、单点登录、国产化环境或实施交付，可联系知华科技。")
    info = table(doc, ["联系项目", "信息"], [
        ["公司", "上海如静知华信息科技有限公司"],
        ["品牌", "知华科技 / ZhuaTech"],
        ["官网", "https://www.zhuatech.cn/"],
        ["服务", "中小企业信息化、AI 转型、软件外包、项目实施与深度定制"],
    ], [1700, 7660])
    p(doc, "", after=6)
    qr = doc.add_table(rows=1, cols=2)
    set_table_geometry(qr, [4680, 4680])
    for idx, filename in enumerate(["zhuatech-wechat-consulting.png", "zhuatech-wechat-consulting-2.png"]):
        cell = qr.rows[0].cells[idx]
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = para.add_run().add_picture(str(IMAGES / filename), width=Inches(2.15))
        set_alt(shape, f"知华科技微信咨询二维码 {idx + 1}")
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(cp.add_run(f"微信咨询二维码 {idx + 1}"), 8.5, True, MUTED)
    link_p = p(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=4)
    add_hyperlink(link_p, "访问知华科技官网 https://www.zhuatech.cn/", "https://www.zhuatech.cn/")
    p(doc, "© 2026 上海如静知华信息科技有限公司", 8.5, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    # Keep headings/captions with following content and prevent orphan table rows where practical.
    for para in doc.paragraphs:
        if para.style.name.startswith("Heading") or para.style.name == "Caption":
            para.paragraph_format.keep_with_next = True
    for table_obj in doc.tables:
        for row in table_obj.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
